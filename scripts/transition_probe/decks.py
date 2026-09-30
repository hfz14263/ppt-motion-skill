"""transition_probe.decks — deck 构造器

**deck 构造** —— 把假设写成 slide XML。纯设置，不需要 PowerPoint
"""

# 本模块由 tools/split_transition_probe.py 从单文件的 build_transition_table.py
# 搬出来（原文照搬，未改逻辑）。对外接口由 __init__.py 重新导出。


from .common import SCRIPTS_DIR, DUR_MS, FPS, GRID_COLS, GRID_ROWS, NS, SLIDE_SECONDS, _load_json
from .data import HYPOTHESES

import io
import json
import os
import re
import zipfile

# --------------------------------------------------------------------------
# The 48 hypotheses. `family` decides how the child is wrapped:
#   core  -> direct child of <p:transition>            (ECMA-376)
#   p14   -> mc:Choice Requires="p14" + fade Fallback  ([MS-PPTX] 3.1)
#   p15   -> mc:Choice Requires="p15" + fade Fallback
#   p159  -> morph, already proven by motion.py
# Attributes are written ONLY where they distinguish two gallery items
# (cube / rotate / orbit are all p14:prism) or where PowerPoint refuses the
# bare form -- 切换/翻转/库/摩天轮/传送带 are refused WITHOUT dir, which the
# roundtrip found and the enum scan explained: PowerPoint itself always writes
# them with dir="l". "Refused" alone would never have said which attribute.
# Everything else is left out ON PURPOSE so the roundtrip reveals PowerPoint's
# own default.
# --------------------------------------------------------------------------
def _with_dir(child, direction):
    """Return `child` with a `dir="..."` attribute injected.

    `dir` goes on the CHILD element (`<p:push dir="l"/>`), not on
    <p:transition> -- that is where PowerPoint itself writes it
    (com-pitfalls §41/§42). Only used by the direction probe; the default
    probe deliberately writes no dir at all.
    """
    if not direction:
        return child
    m = re.match(r"<([A-Za-z0-9:]+)((?:\s[^>]*)?)/>$", child)
    if not m:
        raise ValueError("cannot inject dir into %r" % child)
    tag, attrs = m.group(1), m.group(2) or ""
    attrs = re.sub(r'\s+dir="[^"]*"', "", attrs)      # replace, don't stack
    return '<%s%s dir="%s"/>' % (tag, attrs, direction)


def wrap_transition(family, child, spd="slow", dur=DUR_MS):
    """Return the full XML block for one candidate, per its family.

    MEASURED, not assumed. Two rules came out of the roundtrip:

    * Every transition ends up inside mc:AlternateContent -- even a plain core
      one. The reason is p14:dur: a 2010 attribute needs a Choice/Fallback, and
      PowerPoint wraps the whole <p:transition> rather than dropping dur.
    * The Fallback is NOT always <p:fade/>. For a core element the Fallback
      repeats the core element itself (it is already backwards-compatible);
      only a p14/p15/p159 child degrades to fade. Writing <p:fade/> as the
      Fallback for <p:push/> is accepted but PowerPoint rewrites it -- 16 of
      the first 48 candidates came back "rewritten" purely because of this.
    """
    if family == "core":
        return (
            '<mc:AlternateContent>'
            '<mc:Choice xmlns:p14="%s" Requires="p14">'
            '<p:transition spd="%s" p14:dur="%d">%s</p:transition>'
            '</mc:Choice>'
            '<mc:Fallback><p:transition spd="%s">%s</p:transition>'
            '</mc:Fallback></mc:AlternateContent>'
        ) % (NS["p14"], spd, dur, child, spd, child)
    if family in ("p14", "p15", "p159"):
        prefix = family
        return (
            '<mc:AlternateContent>'
            '<mc:Choice xmlns:%s="%s" Requires="%s">'
            '<p:transition spd="%s" p14:dur="%d">%s</p:transition>'
            '</mc:Choice>'
            '<mc:Fallback><p:transition spd="%s"><p:fade/></p:transition>'
            '</mc:Fallback></mc:AlternateContent>'
        ) % (prefix, NS[prefix], prefix, spd, dur, child, spd)
    raise ValueError("unknown family %r" % family)


def _slide_xml(data, block):
    """Declare every extension namespace on <p:sld> and insert `block` after
    <p:clrMapOvr> (position rule: cSld, clrMapOvr?, transition?, timing? --
    a misplaced block is silently dropped on save, com-pitfalls §12).

    Only the MISSING prefixes are added. python-pptx already writes xmlns:a,
    xmlns:p and xmlns:r on <p:sld>; appending a second xmlns:a makes the part
    not well-formed, and PowerPoint's answer to that is to refuse the WHOLE
    deck (observed: all 48 candidates refused, including a plain <p:fade/>,
    which is how this was found -- §26 again, but the fault was ours)."""
    xml = data.decode("utf-8")
    m = re.search(r"<p:sld(\s[^>]*?)?>", xml)
    have = set(re.findall(r'xmlns:(\w+)=', m.group(1) or ""))
    add = " ".join('xmlns:%s="%s"' % (k, v) for k, v in sorted(NS.items())
                   if k not in have)
    def repl(mm):
        head = mm.group(0)[:-1].rstrip()
        return "%s %s>" % (head, add) if add else mm.group(0)
    xml = re.sub(r"<p:sld(\s[^>]*?)?>", repl, xml, count=1)
    anchor = "</p:clrMapOvr>"
    idx = xml.rindex(anchor) + len(anchor)
    return (xml[:idx] + block + xml[idx:]).encode("utf-8")


def _two_slide_deck(path, label, hyp=None):
    """A minimal deck: slide 1 'from', slide 2 'entering' (+ candidate)."""
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    blank = prs.slide_layouts[6]

    def add(bg, big, small):
        s = prs.slides.add_slide(blank)
        r = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor(*bg)
        r.line.fill.background()
        r.shadow.inherit = False
        tb = s.shapes.add_textbox(Pt(60), Pt(60), prs.slide_width - Pt(120),
                                  prs.slide_height - Pt(120))
        tf = tb.text_frame
        tf.text = big
        p0 = tf.paragraphs[0]
        p0.font.size = Pt(120)
        p0.font.bold = True
        p1 = tf.add_paragraph()
        p1.text = small
        p1.font.size = Pt(20)

    add((235, 220, 200), label, "from")
    add((70, 110, 190), label, hyp["ui"] if hyp else "target")
    prs.save(path)

    if hyp is None:
        return
    buf = io.BytesIO(open(path, "rb").read())
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/slides/slide2\.xml", item.filename):
                data = _slide_xml(data, wrap_transition(hyp["family"],
                                                        hyp["child"]))
            zout.writestr(item, data)
    open(path, "wb").write(out.getvalue())


# --------------------------------------------------------------------------
# anchordeck: does a <p:transition> written on slide N drive N-1 -> N?
# --------------------------------------------------------------------------
# Two claims every user of the table makes, and neither had been measured:
#
#   A. A transition is written on the slide it ANIMATES INTO. So it lives on
#      slide N and drives N-1 -> N. Writing it on slide 1 should do nothing
#      visible. Prove it with rendered frames, not with the XML: the XML is
#      there either way, so "the file contains it" proves nothing.
#   B. <p:transition> takes exactly ONE child. So morph and push cannot both be
#      in force -- they compete for the same slot. (This is what decides whether
#      "shape deformation AND whole-page movement" is even expressible.)
#
# The anchor deck puts an 800ms push on slide 1 and an 800ms wipe on slide 3,
# leaving slide 2 empty. Rendered at 30fps with 2s per slide the boundaries sit
# near frames 60 and 120. If (A) holds there is a multi-frame change burst only
# at the second boundary; the first is a one-frame hard cut.
def _flat_deck(path, labels, blocks, dur=DUR_MS):
    from pptx import Presentation
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    blank = prs.slide_layouts[6]
    palette = [(200, 60, 60), (60, 150, 90), (60, 80, 190), (190, 150, 40)]
    for i, text in enumerate(labels):
        s = prs.slides.add_slide(blank)
        r = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        r.fill.solid()
        r.fill.fore_color.rgb = RGBColor(*palette[i % len(palette)])
        r.line.fill.background()
        r.shadow.inherit = False
        tb = s.shapes.add_textbox(Pt(60), Pt(60), prs.slide_width - Pt(120),
                                  prs.slide_height - Pt(120))
        tf = tb.text_frame
        tf.text = text
        tf.paragraphs[0].font.size = Pt(240)
        tf.paragraphs[0].font.bold = True
    prs.save(path)

    assigned = {}
    buf = io.BytesIO(open(path, "rb").read())
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", item.filename)
            if m and int(m.group(1)) in blocks:
                data = _slide_xml(data, blocks[int(m.group(1))])
                assigned[int(m.group(1))] = blocks[int(m.group(1))]
            zout.writestr(item, data)
    open(path, "wb").write(out.getvalue())
    return assigned


def _shape_deck(path, block, dur=DUR_MS):
    """One 2-slide deck whose 'entering' slide carries `block`.

    Returns the number of slides written so callers can assert nothing was
    silently dropped.
    """
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    blank = prs.slide_layouts[6]
    w, h = prs.slide_width, prs.slide_height

    def cell_colour(i):
        # Identical on both slides: the grid is a STATIC reference frame, not
        # part of the change. An earlier version stepped the palette between
        # the two slides, which made every cell differ -- the diff map went
        # nearly uniform and the direction of the actual motion was averaged
        # away. (push measured as "stationary" because of it.) The only thing
        # that should differ between FROM and TO is the marker below.
        v = (i * 37) % 200 + 30
        return v

    for phase in (0, 1):                      # phase 0 = FROM, 1 = TO
        s = prs.slides.add_slide(blank)
        cw, ch = w // GRID_COLS, h // GRID_ROWS
        for r in range(GRID_ROWS):
            for c in range(GRID_COLS):
                i = r * GRID_COLS + c
                v = cell_colour(i)
                sh = s.shapes.add_shape(1, c * cw, r * ch, cw, ch)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(v, v, v)
                sh.line.fill.background()
                sh.shadow.inherit = False
        # ONE marker that moves and recolours between the two slides. Its
        # displacement makes a directional transition produce a clean, moving
        # diff centroid; its colour change guarantees a strong signal even for
        # effects that do not translate anything (fade, dissolve).
        d = min(w, h) // 3
        dx = 0 if phase == 0 else w // 3       # a known, sizeable shift
        disc = s.shapes.add_shape(9, (w - d) // 2 + dx, (h - d) // 2, d, d)
        disc.fill.solid()
        disc.fill.fore_color.rgb = RGBColor(230, 30, 30) if phase == 0 \
            else RGBColor(30, 30, 230)
        disc.line.fill.background()
        disc.shadow.inherit = False

    prs.save(path)

    buf = io.BytesIO(open(path, "rb").read())
    out = io.BytesIO()
    n = 0
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/slides/slide2\.xml", item.filename):
                data = _slide_xml(data, block)
                n += 1
            zout.writestr(item, data)
    open(path, "wb").write(out.getvalue())
    return n


def _shape_deck2(path, block, dur=DUR_MS):
    """A SECOND probe deck for effects the first one cannot separate.

    Why a second deck exists at all (measured, not theorised): in deck 1 the
    only thing that differs between FROM and TO is a centred disc. Every
    effect that does not TRANSLATE therefore produces the same centre-weighted
    diff, and fade / wipe / dissolve / split / shape / randombar / box all came
    back with radial profiles within 0.03 of each other. They differ in HOW the
    change propagates, not WHERE it sits -- so the probe has to make
    propagation visible.

    Design: FROM and TO are both a fine RANDOM noise texture (different noise on
    each page, same average brightness). Now:
      * a sweeping reveal exposes the new noise region by region -> the diff
        forms a band that TRAVELS, and its direction is measurable;
      * a uniform cross-fade changes every pixel at once -> a flat diff, no
        travel;
      * a split opens from the centre outward -> the diff centroid moves
        outward in both directions.
    Average brightness is matched between the two pages on purpose: if TO were
    brighter overall, a fade would look like a brightness ramp and could be
    mistaken for a sweep.
    """
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.dml.color import RGBColor

    import random as _random

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    blank = prs.slide_layouts[6]
    w, h = prs.slide_width, prs.slide_height
    cols, rows = 48, 27                       # 48x27 tiles = 1296 cells
    tw, th = w // cols, h // rows

    for phase in (0, 1):
        rng = _random.Random(1000 + phase)
        s = prs.slides.add_slide(blank)
        # Two brightness levels only, 50/50: matched mean, but the SPATIAL
        # pattern is what differs. Two levels keep the H.264 noise low.
        for r in range(rows):
            for c in range(cols):
                v = 190 if rng.random() < 0.5 else 60
                sh = s.shapes.add_shape(1, c * tw, r * th, tw, th)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(v, v, v)
                sh.line.fill.background()
                sh.shadow.inherit = False

    prs.save(path)

    buf = io.BytesIO(open(path, "rb").read())
    out = io.BytesIO()
    n = 0
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/slides/slide2\.xml", item.filename):
                data = _slide_xml(data, block)
                n += 1
            zout.writestr(item, data)
    open(path, "wb").write(out.getvalue())
    return n


def _set_attrs(child, overrides):
    """Force (attr, value) pairs onto a child element.

    `value=None` REMOVES the attribute (so the row measures the default form).
    Removal matters: an early version of the dir probe only ever ADDED, so
    "default" rows silently carried whatever the HYPOTHESES child already had.
    """
    m = re.match(r"<([A-Za-z0-9:]+)((?:\s[^>]*)?)/>$", child)
    if not m:
        raise ValueError("cannot set attrs on %r" % child)
    tag, attrs = m.group(1), m.group(2) or ""
    for attr, val in overrides:
        attrs = re.sub(r'\s+' + re.escape(attr) + r'="[^"]*"', "", attrs)
        if val is not None:
            attrs += ' %s="%s"' % (attr, val)
    return "<%s%s/>" % (tag, attrs)


def _probe_deck(out_dir, rows, tag_key, label_of):
    """Shared writer for dirdeck / attrdeck: one deck per probe row."""
    os.makedirs(out_dir, exist_ok=True)
    by_spec = {h["spec"]: h for h in HYPOTHESES}
    written = []
    for spec, label, overrides, why in rows:
        hyp = by_spec.get(spec)
        if hyp is None:
            raise RuntimeError("probe references unknown spec %r" % spec)
        name = "%s_%s_%s" % (spec, tag_key, label)
        path = os.path.join(out_dir, name + ".pptx")
        child = _set_attrs(hyp["child"], overrides)
        block = wrap_transition(hyp["family"], child, dur=DUR_MS)
        n = _shape_deck(path, block)
        if n != 1:
            raise RuntimeError("%s: transition not written (n=%d)" % (name, n))
        entry = {"spec": spec, tag_key: label, "label": label,
                 "ui": hyp["ui"], "group": hyp["group"],
                 "family": hyp["family"], "child": child, "why": why,
                 "deck": name}
        written.append(entry)
        json.dump({"case": "%s probe: %s %s=%s" % (tag_key, spec, tag_key,
                                                   label),
                   "spec": spec, tag_key: label, "label": label,
                   "ui": hyp["ui"], "group": hyp["group"],
                   "family": hyp["family"], "child": child,
                   "dur_ms": DUR_MS, "fps": FPS, "slides": 2,
                   "transition_on": 2,
                   "expected_boundary_frames": [SLIDE_SECONDS * FPS],
                   "grid": [GRID_COLS, GRID_ROWS]},
                  io.open(os.path.join(out_dir, name + ".manifest.json"),
                          "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump({"decks": written, "fps": FPS, "dur_ms": DUR_MS,
               "slide_seconds": SLIDE_SECONDS, "count": len(written)},
              io.open(os.path.join(out_dir, "_index.json"), "w",
                      encoding="utf-8"), ensure_ascii=False, indent=1)
    return written


def _anim_deck(path, anim, trans_block=None):
    """One deck whose 'entering' slide carries a REAL <p:timing> animation and
    optionally a <p:transition>.

    The timing block is generated by motion.py's own build_timing(), not
    hand-written here. That matters: the question is whether the SHIPPING
    engine's output coexists with a transition, and a hand-rolled timing block
    would answer a question nobody asked.

    The animation targets whatever shape the probe deck already has, by
    position, so the deck geometry stays the shape-probe geometry (the marker
    disc is what moves -- a static grid would make a fade invisible).

    Returns (n_timing, n_transition): how many slides got each block, so the
    caller can assert nothing was silently dropped.
    """
    from pptx import Presentation
    from pptx.util import Emu
    from pptx.dml.color import RGBColor
    import motion as _motion

    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    blank = prs.slide_layouts[6]
    w, h = prs.slide_width, prs.slide_height

    def cell_colour(i):
        v = (i * 37) % 200 + 30
        return v

    for phase in (0, 1):
        s = prs.slides.add_slide(blank)
        cw, ch = w // GRID_COLS, h // GRID_ROWS
        for r in range(GRID_ROWS):
            for c in range(GRID_COLS):
                i = r * GRID_COLS + c
                v = cell_colour(i)
                sh = s.shapes.add_shape(1, c * cw, r * ch, cw, ch)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(v, v, v)
                sh.line.fill.background()
                sh.shadow.inherit = False
        d = min(w, h) // 3
        dx = 0 if phase == 0 else w // 3
        disc = s.shapes.add_shape(9, (w - d) // 2 + dx, (h - d) // 2, d, d)
        disc.fill.solid()
        disc.fill.fore_color.rgb = RGBColor(230, 30, 30) if phase == 0 \
            else RGBColor(30, 30, 230)
        disc.line.fill.background()
        disc.shadow.inherit = False

    prs.save(path)

    # Find the moving marker's shape id on slide 2, so the animation has a real
    # target (motion.py resolves targets by name/id from the spTree).
    buf = io.BytesIO(open(path, "rb").read())
    with zipfile.ZipFile(buf) as z:
        s2 = z.read("ppt/slides/slide2.xml").decode("utf-8")
    ids = re.findall(r'<p:cNvPr id="(\d+)"[^>]*name="([^"]*)"', s2)
    if not ids:
        raise RuntimeError("no shapes found on slide 2 of %s" % path)
    # The LAST shape is the disc (grid is written first).
    marker_id = ids[-1][0]

    timing_block = ""
    if anim is not None:
        catalog = _load_json(os.path.join(SCRIPTS_DIR, "motion_catalog.json"))
        meta = catalog.get(str(anim["effect"]))
        if meta is None:
            raise RuntimeError("unknown anim effect %r" % anim["effect"])
        entries = [{"shape_id": marker_id, "tpl": meta["template"],
                    "effect": anim, "delay": anim.get("delay") or 0.0}]
        timing_block = _motion.build_timing(entries, _motion.NodeIds(4))

    n_t, n_x = 0, 0
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r"ppt/slides/slide2\.xml", item.filename):
                xml = data.decode("utf-8")
                # Declare extension prefixes on <p:sld> (only the missing ones:
                # a duplicate makes the part malformed, com-pitfalls §26/§12).
                m = re.search(r"<p:sld(\s[^>]*?)?>", xml)
                have = set(re.findall(r'xmlns:(\w+)=', m.group(1) or ""))
                add = " ".join('xmlns:%s="%s"' % (k, v)
                               for k, v in sorted(NS.items()) if k not in have)
                if add:
                    xml = re.sub(r"<p:sld(\s[^>]*?)?>",
                                 lambda mm: "%s %s>" % (mm.group(0)[:-1].rstrip(), add),
                                 xml, count=1)
                # ORDER: cSld, clrMapOvr?, transition?, timing?, extLst?
                # Insert by walking the anchor forward, never by re-deriving an
                # index from string searches -- that is exactly how a timing
                # block ends up INSIDE the transition (com-pitfalls §12/§19).
                anchor = "</p:clrMapOvr>"
                idx = xml.rindex(anchor) + len(anchor)
                pieces = []
                if trans_block:
                    pieces.append(trans_block)
                    n_x += 1
                if timing_block:
                    pieces.append(timing_block)
                    n_t += 1
                if pieces:
                    xml = xml[:idx] + "".join(pieces) + xml[idx:]
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    open(path, "wb").write(out.getvalue())
    return n_t, n_x
