# -*- coding: utf-8 -*-
"""构造"同图双版本"版式的样例页，供 verify_dual_photo.py 做正/负对照。

版式定义（来自抖音教学视频复现）：
    同一张照片，两个版本 —— 淡化(可选虚化)整图铺底 + 原色塞进形状当前景。
    只有两层。视频里"复制+翻转+重叠"是过程演示形状，成片前必须删掉。

用法:
    python build_dual_photo.py --src 人像.png --out ok.pptx
    python build_dual_photo.py --src 人像.png --out fat.pptx  --stretch 1.35
    python build_dual_photo.py --src 人像.png --out diff.pptx --bg-src 另一张.png
    python build_dual_photo.py --src 人像.png --out demo.pptx --shape chevron

--stretch > 1 会故意把人物横向拉胖（复现视频里的失败对照"拉伸人物就长胖"），
用来验证校验器能量化出这个失真倍数。
"""
import argparse
import os
import shutil
import tempfile

from lxml import etree
from PIL import Image, ImageEnhance, ImageOps, ImageFilter
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# 画布沿用 dsh-ppt-studio 约定：960x540，1px=1pt
SW, SH = 960.0, 540.0
# 主形状默认位置（与视频复现成品一致）
DEFAULTS = {"left": 317.0, "top": 15.0, "width": 557.0, "height": 505.0}

SHAPES = {
    "chevron": MSO_SHAPE.CHEVRON,
    "roundrect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rect": MSO_SHAPE.RECTANGLE,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
}


def wash(src, dst, brightness=2.4, contrast=0.66, blur=6.0,
         black=(70, 50, 34), white=(243, 227, 205)):
    """淡化版：去色 + 提亮压对比 + 上暖棕色调 + 轻微虚化。"""
    im = Image.open(src)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        flat = Image.new("RGB", im.size, (255, 255, 255))
        flat.paste(im, mask=im.split()[-1])
        im = flat
    else:
        im = im.convert("RGB")
    g = ImageEnhance.Color(im).enhance(0.0)
    g = ImageEnhance.Brightness(g).enhance(brightness)
    g = ImageEnhance.Contrast(g).enhance(contrast)
    # colorize 只吃 8-bit 灰度，Color.enhance(0) 的产物仍是 RGB，必须显式转 L
    g = ImageOps.colorize(g.convert("L"), black=black, white=white)
    if blur:
        g = g.filter(ImageFilter.GaussianBlur(blur))
    g.save(dst, quality=92)
    return g


def _crop_for(src_ar, dst_ar, stretch=1.0):
    """算出让填充不变形的 srcRect 裁剪比例 (kw, kh)。
    stretch>1 时故意少取横向内容 -> 人物被横向拉胖，失真倍数 = stretch。"""
    if src_ar >= dst_ar:                    # 源图偏宽：裁左右
        kw = dst_ar / src_ar
        kh = 1.0
    else:                                   # 源图偏高：裁上下
        kw = 1.0
        kh = src_ar / dst_ar
    if stretch >= 1.0:
        kw = kw / stretch
    else:
        kh = kh * stretch
    if not (0 < kw <= 1.0) or not (0 < kh <= 1.0):
        raise ValueError("stretch=%.3f 超出可表达范围 (kw=%.3f, kh=%.3f)" % (stretch, kw, kh))
    return kw, kh


def picture_fill(shape, slide, img_path, src_rect=None):
    """手写 OOXML blipFill（python-pptx 没有形状图片填充 API）。
    坑位：spPr 内 prstGeom 之后、ln 之前；先删旧 fill，否则 xsd:choice 取第一个。"""
    part, rId = slide.part.get_or_add_image_part(img_path)
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    bf = etree.Element(qn("a:blipFill"))
    bf.set("rotWithShape", "1")
    blip = etree.SubElement(bf, qn("a:blip"))
    blip.set(qn("r:embed"), rId)
    if src_rect:
        sr = etree.SubElement(bf, qn("a:srcRect"))
        for k, v in zip(("l", "t", "r", "b"), src_rect):
            if v:
                sr.set(k, str(v))
    st = etree.SubElement(bf, qn("a:stretch"))
    etree.SubElement(st, qn("a:fillRect"))
    anchor = None
    for tag in ("a:ln", "a:effectLst", "a:scene3d", "a:sp3d", "a:extLst"):
        f = spPr.find(qn(tag))
        if f is not None:
            anchor = f
            break
    if anchor is not None:
        anchor.addprevious(bf)
    else:
        g = spPr.find(qn("a:prstGeom"))
        if g is None:
            g = spPr.find(qn("a:custGeom"))
        g.addnext(bf)
    shape.line.fill.background()
    shape.shadow.inherit = False


def build(src, out, shape="chevron", stretch=1.0, bg_src=None,
          geom=None, blur=6.0, tmpdir=None):
    """造一页"同图双版本"。返回 dict（含实际失真倍数 ratio，供测试核对）。"""
    geom = geom or DEFAULTS
    tmp = tmpdir or tempfile.mkdtemp(prefix="dualphoto_")
    bg_path = os.path.join(tmp, "bg_washed.jpg")
    wash(bg_src or src, bg_path, blur=blur)

    prs = Presentation()
    # python-pptx 默认模板的作者是它自己的作者名，会被 privacy_audit 判为"人名泄漏"。
    # 项目约定：这两个字段写工具名 dsh-ppt-studio（见 tests/privacy_audit.py AUTHOR_ALLOW）。
    prs.core_properties.author = "dsh-ppt-studio"
    prs.core_properties.last_modified_by = "dsh-ppt-studio"
    prs.slide_width = Pt(SW)
    prs.slide_height = Pt(SH)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1) 淡化整图铺底（裁长边保比例）
    src_im = Image.open(bg_src or src)
    ar_pic = src_im.width / src_im.height
    pic = slide.shapes.add_picture(bg_path, 0, 0, Pt(SW), Pt(SH))
    ar = SW / SH
    if ar_pic > ar:
        keep = ar / ar_pic
        pic.crop_left = pic.crop_right = (1 - keep) / 2
    else:
        keep = ar_pic / ar
        pic.crop_top = pic.crop_bottom = (1 - keep) / 2

    # 2) 原色塞进形状
    im = Image.open(src)
    src_ar = im.width / im.height
    dst_ar = geom["width"] / geom["height"]
    kw, kh = _crop_for(src_ar, dst_ar, stretch)
    # 千分百分比*100（100000 = 100%）
    l = r = int(round((1 - kw) / 2 * 100000))
    t = b = int(round((1 - kh) / 2 * 100000))
    shp = slide.shapes.add_shape(
        SHAPES[shape], Pt(geom["left"]), Pt(geom["top"]),
        Pt(geom["width"]), Pt(geom["height"]))
    picture_fill(shp, slide, src, src_rect=(l, t, r, b))

    prs.save(out)
    # 校验器应当量出的失真倍数
    ratio = dst_ar / (src_ar * kw / kh)
    return {"out": out, "stretch_requested": stretch, "ratio": ratio,
            "src_rect": (l, t, r, b), "tmp": tmp}


def main():
    ap = argparse.ArgumentParser(description="构造同图双版本版式样例")
    ap.add_argument("--src", required=True, help="前景原图（也是默认的背景原图）")
    ap.add_argument("--out", required=True)
    ap.add_argument("--bg-src", default=None, help="背景用另一张图（造'非同图'负对照）")
    ap.add_argument("--shape", default="chevron", choices=sorted(SHAPES))
    ap.add_argument("--stretch", type=float, default=1.0, help=">1 故意横向拉胖")
    ap.add_argument("--blur", type=float, default=6.0)
    a = ap.parse_args()
    info = build(a.src, a.out, shape=a.shape, stretch=a.stretch,
                 bg_src=a.bg_src, blur=a.blur)
    print("saved: %s" % info["out"])
    print("srcRect(l,t,r,b) = %s" % (info["src_rect"],))
    print("ratio(校验器应量出) = %.4f  (请求 stretch = %.4f)" % (info["ratio"], a.stretch))


if __name__ == "__main__":
    main()
